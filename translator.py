import logging
import os
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import List, Tuple, Dict, Optional, Generator

import openpyxl
from dotenv import load_dotenv
from langchain_openai import ChatOpenAI
from pydantic import BaseModel
from pricing import MODEL_PRICING

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s %(levelname)s %(message)s',
    handlers=[logging.StreamHandler()]
)
logger = logging.getLogger(__name__)

# Utility to suppress info logs for CLI usage
def suppress_info_logging():
    logger.setLevel(logging.WARNING)

# Load environment variables from .env at the start
load_dotenv()
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")


class TranslationOutput(BaseModel):
    translations: List[str]


class RateLimiter:
    def __init__(self, max_calls_per_minute):
        self.period = 60.0 / max_calls_per_minute
        self.lock = threading.Lock()
        self.last_call = 0.0

    def wait(self):
        with self.lock:
            now = time.time()
            wait_time = self.last_call + self.period - now
            if wait_time > 0:
                time.sleep(wait_time)
            self.last_call = max(now, self.last_call + self.period)


class XLSXTranslator:
    def __init__(self, input_path: str, target_language: str = "en", max_workers: int = 11, model_name: str = "gpt-4o",
                 rpm_limit: int = 500):
        self.input_path = input_path
        self.target_language = target_language
        self.max_workers = max_workers
        self.model_name = model_name
        self.rpm_limit = rpm_limit
        self.text_cells: List[Tuple[str, str, str]] = []
        self.texts: List[str] = []
        self.sheet_names: List[str] = []
        self.all_translations: List[Optional[str]] = []
        self.sheet_name_translations: List[Optional[str]] = []
        self.progress = 0
        self.done = False
        self.error: Optional[str] = None
        self._llm = None
        self.actual_input_tokens = 0
        self.actual_output_tokens = 0
        self.actual_cost = 0.0
        self._prepare()

    def _prepare(self):
        wb = openpyxl.load_workbook(self.input_path)
        self.sheet_names = [sheet.title for sheet in wb.worksheets if getattr(sheet, 'sheet_state', 'visible') == 'visible']
        self.text_cells = self.extract_text_locations(wb)
        self.texts = [v for _, _, v in self.text_cells]
        self.all_translations = [None] * len(self.texts)
        self.sheet_name_translations = [None] * len(self.sheet_names)
        self.progress = 0
        self.done = False
        self.error = None
        self._llm = ChatOpenAI(model=self.model_name, openai_api_key=OPENAI_API_KEY, temperature=0, max_retries=23)
        logger.info(
            f"Prepared XLSXTranslator for file '{self.input_path}' to '{self.target_language}' using model '{self.model_name}' with {self.max_workers} threads.")

    @staticmethod
    def extract_text_locations(wb) -> List[Tuple[str, str, str]]:
        text_cells = []  # (sheet_name, cell, value)
        for sheet in wb.worksheets:
            if getattr(sheet, 'sheet_state', 'visible') != 'visible':
                continue
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        val = cell.value.strip()
                        # Only include if at least one alphanumeric character (unicode-aware)
                        if val and any(char.isalnum() for char in val):
                            text_cells.append((sheet.title, cell.coordinate, cell.value))
        return text_cells

    @staticmethod
    def sanitize_sheet_title(title: str, used_titles: set) -> str:
        # Excel sheet name rules: max 31 chars, cannot contain : \ / ? * [ ]
        invalid_chars = r'[:\\/?*\[\]]'
        sanitized = re.sub(invalid_chars, '', title)
        sanitized = sanitized.strip()
        if not sanitized:
            sanitized = 'Sheet'
        sanitized = sanitized[:31]
        # Ensure uniqueness
        base = sanitized
        i = 1
        while sanitized in used_titles:
            suffix = f"_{i}"
            sanitized = (base[:31 - len(suffix)] + suffix) if len(base) + len(suffix) > 31 else base + suffix
            i += 1
        used_titles.add(sanitized)
        return sanitized

    @staticmethod
    def extract_sheet_references_from_formula(formula: str) -> set:
        """
        Extract all sheet references from a formula string.
        Returns a set of sheet names as they appear in the formula.
        """
        # Matches Sheet! or 'Sheet Name'! before a cell reference
        pattern = r"(?:'([^']+)'|([A-Za-z0-9_]+))!"
        refs = set()
        for m in re.finditer(pattern, formula):
            sheet = m.group(1) if m.group(1) is not None else m.group(2)
            refs.add(sheet)
        return refs

    @staticmethod
    def adjust_formula_sheet_references(formula: str, sheet_name_map: Dict[str, str]) -> str:
        """
        Replaces sheet references in a formula string using the provided mapping.
        Handles cases like Sheet1!A1, 'My Sheet'!A1, etc.
        """
        def repl(match):
            sheet = match.group(1)
            # Remove quotes if present
            sheet_unquoted = sheet[1:-1] if sheet.startswith("'") and sheet.endswith("'") else sheet
            new_sheet = sheet_name_map.get(sheet_unquoted, sheet_unquoted)
            # Add quotes if new sheet name contains spaces or special chars
            if re.search(r"[\s:\\/?*\[\]']", new_sheet):
                new_sheet = f"'{new_sheet}'"
            return f"{new_sheet}!"

        # Regex: matches Sheet! or 'My Sheet'! after '=' or any operator
        # Handles formulas like =SheetA!A1+1, =SUM(SheetA!A1,SheetB!B2)
        pattern = re.compile(r"('([^']+)'|" + '|'.join(re.escape(name) for name in sheet_name_map.keys()) + ")!")
        result = re.sub(pattern, repl, formula)
        # Debug print
        print(f"DEBUG adjust_formula_sheet_references: input={formula}, output={result}")
        return result

    @staticmethod
    def adjust_formula_sheet_references_force(formula: str, orig_to_final: Dict[str, str], orig_to_translated: Dict[str, str]) -> str:
        """
        Forcibly replace all sheet references in a formula, using the mapping from original sheet names to final translated sheet names.
        orig_to_final: {original sheet name: final sanitized translated name}
        orig_to_translated: {original sheet name: LLM translated name}
        """
        def repl(match):
            sheet_ref = match.group(1) if match.group(1) is not None else match.group(2)
            # Find which original sheet this translated reference corresponds to
            for orig, translated in orig_to_translated.items():
                if translated == sheet_ref:
                    final = orig_to_final.get(orig, translated)
                    # Add quotes if needed
                    if re.search(r"[\s:\\/?*\[\]']", final):
                        final = f"'{final}'"
                    return f"{final}!"
            # If not found, leave as is
            return match.group(0)
        pattern = r"(?:'([^']+)'|([A-Za-z0-9_ ]+))!"
        return re.sub(pattern, repl, formula)

    @staticmethod
    def adjust_formula_sheet_references_force_precise(formula: str, orig_to_final: Dict[str, str]) -> str:
        """
        Forcibly replace all sheet references in a formula, using the mapping from original sheet names to final tab names.
        Always quotes the tab name if it contains spaces or special characters.
        """
        def repl(match):
            sheet_ref = match.group(1) if match.group(1) is not None else match.group(2)
            # Always use the actual tab name for this reference
            final = orig_to_final.get(sheet_ref, sheet_ref)
            # Add quotes if needed
            if re.search(r"[\s:\\/?*\[\]']", final):
                final = f"'{final}'"
            return f"{final}!"
        pattern = r"(?:'([^']+)'|([A-Za-z0-9_ ]+))!"
        return re.sub(pattern, repl, formula)

    @staticmethod
    def update_formulas_with_translated_sheetnames(wb, sheet_name_map):
        """
        For every formula in visible sheets, replace original sheet name references with translated names.
        """
        # Build a regex pattern for all original sheet names (longest first to avoid partial matches)
        orig_names = sorted(sheet_name_map.keys(), key=lambda x: -len(x))
        # Pattern matches: 'Sheet Name'! or SheetName!
        pattern = re.compile(r"('([^']+)'|" + '|'.join(re.escape(name) for name in orig_names) + ")!")
        visible_sheets = [sheet for sheet in wb.worksheets if getattr(sheet, 'sheet_state', 'visible') == 'visible']
        for ws in visible_sheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith('='):
                        def repl(match):
                            sheet_ref = match.group(2) if match.group(2) is not None else match.group(1)
                            # Remove quotes if present
                            sheet_unquoted = sheet_ref[1:-1] if sheet_ref and sheet_ref.startswith("'") and sheet_ref.endswith("'") else sheet_ref
                            translated = sheet_name_map.get(sheet_unquoted, sheet_unquoted)
                            # Quote if needed
                            if re.search(r"[\s:\\/?*\[\]']", translated):
                                translated = f"'{translated}'"
                            return f"{translated}!"
                        cell.value = pattern.sub(repl, cell.value)

    @staticmethod
    def replace_texts_and_sheet_names_in_xlsx(xlsx_path: str, replacements: Dict[Tuple[str, str], str],
                                              sheet_name_map: Dict[str, str], output_path: str):
        wb = openpyxl.load_workbook(xlsx_path)
        # Only operate on visible sheets
        visible_sheets = [sheet for sheet in wb.worksheets if getattr(sheet, 'sheet_state', 'visible') == 'visible']
        visible_sheet_names = [sheet.title for sheet in visible_sheets]
        # Step 1: Build sanitized mapping for sheet names
        used_titles = set(visible_sheet_names)
        sanitized_map = {}
        for old_name, new_name in sheet_name_map.items():
            if old_name not in visible_sheet_names:
                continue
            sanitized = XLSXTranslator.sanitize_sheet_title(new_name, used_titles)
            sanitized_map[old_name] = sanitized
        # Step 2: Replace formulas with translated sheet names
        XLSXTranslator.update_formulas_with_translated_sheetnames(wb, sanitized_map)
        # Step 3: Replace cell values (non-formulas) as before
        for (sheet_name, cell_coord), new_value in replacements.items():
            if sheet_name not in visible_sheet_names:
                continue
            ws = wb[sheet_name]
            if not (isinstance(new_value, str) and new_value.startswith('=')):
                ws[cell_coord].value = new_value
        # Step 4: Rename sheets
        for old_name, sanitized in sanitized_map.items():
            if old_name != sanitized:
                ws = wb[old_name]
                ws.title = sanitized
        wb.save(output_path)

    def _translate_single(self, text: str) -> str:
        prompt = (
            f"You are a translation engine. Translate the following text to {self.target_language}. "
            f"Return only the translation with no extra text.\n"
            f"You must produce translations line by line, one original line of text corresponds one line of translation.\n"
            f"Input: {text}"
        )
        try:
            result = self._llm.invoke(prompt)
            # Extract token usage and sum
            usage = getattr(result, "response_metadata", {}).get("token_usage")
            if usage:
                price_info = MODEL_PRICING.get(self.model_name, MODEL_PRICING.get("gpt-4o"))
                self.actual_cost += (usage.get("prompt_tokens", 0) / 1000 * price_info["input"]) + (usage.get("completion_tokens", 0) / 1000 * price_info["output"])
                self.actual_input_tokens += usage.get("prompt_tokens", 0)
                self.actual_output_tokens += usage.get("completion_tokens", 0)
            if logger.isEnabledFor(logging.INFO):
                logger.info(f"Translated: '{text[:30]}...' -> '{result.content.strip()[:30]}...'")
            return result.content.strip()
        except Exception as e:
            logger.error(f"Translation failed for: '{text[:30]}...': {e}")
            return f"[ERROR: {e}]"

    def _translate_single_threadsafe(self, text: str, rate_limiter: 'RateLimiter') -> str:
        rate_limiter.wait()
        return self._translate_single(text)

    def translate(self):
        try:
            logger.info(f"Starting translation of {len(self.texts)} cells and {len(self.sheet_names)} sheet names.")
            total = len(self.texts) + len(self.sheet_names)
            completed = 0
            rate_limiter = RateLimiter(self.rpm_limit)
            # Translate cell values concurrently, but rate-limited
            with ThreadPoolExecutor(max_workers=self.max_workers) as executor:
                futures = {executor.submit(self._translate_single_threadsafe, text, rate_limiter): idx for idx, text in
                           enumerate(self.texts)}
                for future in as_completed(futures):
                    idx = futures[future]
                    try:
                        self.all_translations[idx] = future.result()
                    except Exception as e:
                        self.all_translations[idx] = f"[ERROR: {e}]"
                        logger.error(f"Cell translation error: {e}")
                    completed += 1
                    self.progress = completed / total
            # Translate sheet names concurrently, but rate-limited
            with ThreadPoolExecutor(max_workers=self.max_workers) as executor:
                futures = {executor.submit(self._translate_single_threadsafe, name, rate_limiter): idx for idx, name in
                           enumerate(self.sheet_names)}
                for future in as_completed(futures):
                    idx = futures[future]
                    try:
                        self.sheet_name_translations[idx] = future.result()
                    except Exception as e:
                        self.sheet_name_translations[idx] = f"[ERROR: {e}]"
                        logger.error(f"Sheet name translation error: {e}")
                    completed += 1
                    self.progress = completed / total
            self.done = True
            logger.info("Translation completed.")
        except Exception as e:
            self.error = str(e)
            self.done = True
            logger.error(f"Fatal error during translation: {e}")

    def translate_with_progress(self) -> Generator[float, None, None]:
        try:
            lock = threading.Lock()
            completed = 0
            total = len(self.texts) + len(self.sheet_names)
            rate_limiter = RateLimiter(self.rpm_limit)
            # Translate cell values concurrently, but rate-limited
            with ThreadPoolExecutor(max_workers=self.max_workers) as executor:
                futures = {executor.submit(self._translate_single_threadsafe, text, rate_limiter): idx for idx, text in
                           enumerate(self.texts)}
                for future in as_completed(futures):
                    idx = futures[future]
                    try:
                        self.all_translations[idx] = future.result()
                    except Exception as e:
                        self.all_translations[idx] = f"[ERROR: {e}]"
                    completed += 1
                    with lock:
                        self.progress = completed / total
                    yield self.progress
            # Translate sheet names concurrently, but rate-limited
            with ThreadPoolExecutor(max_workers=self.max_workers) as executor:
                futures = {executor.submit(self._translate_single_threadsafe, name, rate_limiter): idx for idx, name in
                           enumerate(self.sheet_names)}
                for future in as_completed(futures):
                    idx = futures[future]
                    try:
                        self.sheet_name_translations[idx] = future.result()
                    except Exception as e:
                        self.sheet_name_translations[idx] = f"[ERROR: {e}]"
                    completed += 1
                    with lock:
                        self.progress = completed / total
                    yield self.progress
            self.done = True
        except Exception as e:
            self.error = str(e)
            self.done = True
            yield 1.0

    def get_progress(self) -> float:
        return self.progress

    def get_result(self, output_path: str) -> Optional[str]:
        if self.done and not self.error:
            replacements = {(sheet, cell): t for (sheet, cell, _), t in zip(self.text_cells, self.all_translations)}
            sheet_name_map = {old: new for old, new in zip(self.sheet_names, self.sheet_name_translations)}
            self.replace_texts_and_sheet_names_in_xlsx(self.input_path, replacements, sheet_name_map, output_path)
            logger.info(f"Wrote translated XLSX to '{output_path}'")
            return output_path
        logger.error(f"Translation failed or incomplete; output not written.")
        return None


def translate_xlsx_file(input_path: str, output_path: str, target_language: str = "en",
                        model_name: str = "gpt-4o") -> None:
    """
    Translates all text in an XLSX file and writes the result to output_path.
    """
    translator = XLSXTranslator(input_path, target_language, model_name=model_name)
    translator.translate()
    translator.get_result(output_path)
