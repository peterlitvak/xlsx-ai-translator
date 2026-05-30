import unittest
import os
import tempfile

import openpyxl

from translator import XLSXTranslator

test_input = "test_input.xlsx"
test_output = "test_output.xlsx"
OPENAI_TEST_MODEL = os.getenv("OPENAI_TEST_MODEL", "gpt-4o-mini")


class TestTranslator(unittest.TestCase):
    def setUp(self):
        # Create a simple xlsx file
        self.test_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Sheet1"
        ws["A1"] = "Hello"
        ws["B2"] = "World"
        wb.save(self.test_file.name)
        self.output_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")

    def tearDown(self):
        os.remove(self.test_file.name)
        if os.path.exists(self.output_file.name):
            os.remove(self.output_file.name)
        for f in [test_input, test_output]:
            if os.path.exists(f):
                os.remove(f)

    def test_real_llm_translation_flow(self):
        translator = XLSXTranslator(
            self.test_file.name,
            target_language="fr",
            model_name=OPENAI_TEST_MODEL,
            max_workers=1,
        )
        translator.translate()
        self.assertIsNone(translator.error)
        result_path = translator.get_result(self.output_file.name)
        self.assertIsNotNone(result_path)
        assert result_path is not None
        wb = openpyxl.load_workbook(result_path)
        visible_values = [
            str(cell.value)
            for sheet in wb.worksheets
            if getattr(sheet, "sheet_state", "visible") == "visible"
            for row in sheet.iter_rows()
            for cell in row
            if cell.value
        ]
        self.assertGreaterEqual(len(visible_values), 2)
        print("Translated values:", visible_values)

    def test_formula_sheet_reference_adjustment(self):
        # Create workbook with two sheets and a formula referencing the other sheet
        wb = openpyxl.Workbook()
        ws1 = wb.active
        assert ws1 is not None
        ws1.title = "SheetA"
        ws2 = wb.create_sheet(title="SheetB")
        ws1["A1"] = 42
        ws2["B2"] = "=SheetA!A1+1"
        temp_input = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        temp_output = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        wb.save(temp_input.name)
        # Prepare replacements and mapping
        replacements = {("SheetB", "B2"): "=SheetA!A1+1"}
        sheet_name_map = {"SheetA": "TranslatedA", "SheetB": "TranslatedB"}
        # Run replacement
        XLSXTranslator.replace_texts_and_sheet_names_in_xlsx(
            temp_input.name, replacements, sheet_name_map, temp_output.name
        )
        # Check result
        wb2 = openpyxl.load_workbook(temp_output.name)
        ws2_new = wb2["TranslatedB"]
        formula = ws2_new["B2"].value
        print("DEBUG: formula in output:", formula)
        self.assertTrue(
            formula.startswith("=TranslatedA!A1+1")
            or formula.startswith("='TranslatedA'!A1+1")
        )
        os.remove(temp_input.name)
        os.remove(temp_output.name)

    def test_real_translation_with_formula_and_sheet_name(self):
        # Create workbook with two sheets and a formula referencing the other sheet
        wb = openpyxl.Workbook()
        ws1 = wb.active
        assert ws1 is not None
        ws1.title = "Data"
        ws2 = wb.create_sheet(title="Summary")
        ws1["A1"] = "Total"
        ws1["B1"] = 100
        ws2["A1"] = "=Data!B1"
        temp_input = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        temp_output = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        wb.save(temp_input.name)
        # Run translation (to French for example)
        translator = XLSXTranslator(
            temp_input.name,
            target_language="fr",
            model_name=OPENAI_TEST_MODEL,
            max_workers=1,
        )
        translator.translate()
        result_path = translator.get_result(temp_output.name)
        self.assertIsNotNone(result_path)
        assert result_path is not None
        wb2 = openpyxl.load_workbook(result_path)
        visible_sheet_names = {
            sheet.title
            for sheet in wb2.worksheets
            if getattr(sheet, "sheet_state", "visible") == "visible"
        }
        formulas = [
            cell.value
            for sheet in wb2.worksheets
            if getattr(sheet, "sheet_state", "visible") == "visible"
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        self.assertTrue(formulas)
        for formula in formulas:
            self.assertTrue(
                any(sheet_name in formula for sheet_name in visible_sheet_names),
                f"Formula {formula} does not reference a visible translated sheet.",
            )
        os.remove(temp_input.name)
        os.remove(temp_output.name)

    def test_japanese_sample_file_formula_integrity(self):
        # This test assumes test.xlsx exists in the project directory and contains formulas and Japanese text
        import tempfile

        src = os.path.join(os.path.dirname(__file__), "test.xlsx")
        temp_output = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        # Translate to English (or any language)
        translator = XLSXTranslator(
            src,
            target_language="en",
            model_name=OPENAI_TEST_MODEL,
            max_workers=1,
        )
        translator.translate()
        result_path = translator.get_result(temp_output.name)
        self.assertIsNotNone(result_path)
        assert result_path is not None
        wb = openpyxl.load_workbook(result_path)
        # Only consider visible sheets
        visible_sheets = [
            sheet
            for sheet in wb.worksheets
            if getattr(sheet, "sheet_state", "visible") == "visible"
        ]
        visible_sheet_names = [sheet.title for sheet in visible_sheets]
        # Collect all formulas and check that their sheet references exist in the translated file
        for ws in visible_sheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        # Extract referenced sheets from formula
                        refs = XLSXTranslator.extract_sheet_references_from_formula(
                            cell.value
                        )
                        for ref in refs:
                            # Remove quotes if present
                            ref_clean = (
                                ref[1:-1]
                                if ref.startswith("'") and ref.endswith("'")
                                else ref
                            )
                            self.assertIn(
                                ref_clean,
                                visible_sheet_names,
                                f"Formula {cell.value} in sheet {ws.title} references missing sheet {ref_clean}",
                            )
                        print(
                            f"Sheet: {ws.title}, Cell: {cell.coordinate}, Formula: {cell.value}"
                        )
        temp_output.close()
        os.remove(temp_output.name)


if __name__ == "__main__":
    unittest.main()
