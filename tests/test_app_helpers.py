import os
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from typing import ClassVar, Iterable, Optional
from unittest.mock import patch

import openpyxl

from models.translation_usage import TranslationUsage
from services.actual_usage_state import (
    ACTUAL_COST_KEY,
    ACTUAL_INPUT_TOKENS_KEY,
    ACTUAL_OUTPUT_TOKENS_KEY,
    ACTUAL_USAGE_SOURCE_KEY,
    LEGACY_SHOW_ACTUAL_USAGE_KEY,
    clear_actual_usage_on_source_change,
    get_actual_usage_for_source,
    store_actual_usage,
)
from services.translation_estimator import (
    estimate_output_tokens,
    estimate_translation_costs,
    estimate_xlsx_file,
    extract_translatable_texts,
)
from services.download_artifacts import (
    DOWNLOAD_ARTIFACT_DIR_KEY,
    TMPDIR_ENV,
    TRANSLATION_WORK_DIR_ENV,
    cleanup_download_artifact,
    ensure_translation_work_dirs,
    stage_download_artifact,
)
from services.translation_workflow import (
    TranslationError,
    ZipTranslationProgress,
    translate_single_xlsx,
    translate_xlsx_zip,
    XLSX_MIME_TYPE,
    ZIP_MIME_TYPE,
)
from utils.file_names import (
    is_zip_filename,
    is_xlsx_filename,
    translated_zip_name,
    translated_xlsx_name,
)
from utils.zip_archives import (
    UnsafeZipError,
    find_xlsx_files,
    safe_extract_zip,
)


class TestAppHelperUnitTests(unittest.TestCase):
    def create_workbook_buffer(self) -> BytesIO:
        """Create an in-memory workbook with visible and hidden sheet text."""
        workbook = openpyxl.Workbook()
        visible_sheet = workbook.active
        assert visible_sheet is not None
        visible_sheet.title = "Visible"
        visible_sheet["A1"] = "Hello"
        visible_sheet["A2"] = "   "
        visible_sheet["A3"] = "!!!"
        visible_sheet["A4"] = 123

        hidden_sheet = workbook.create_sheet("Hidden")
        hidden_sheet.sheet_state = "hidden"
        hidden_sheet["A1"] = "Secret"

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    def test_extract_translatable_texts_ignores_hidden_sheets_and_non_translatable_values(
        self,
    ) -> None:
        workbook_buffer = self.create_workbook_buffer()

        texts = extract_translatable_texts(workbook_buffer)

        self.assertEqual(["Hello"], texts)

    def test_estimate_output_tokens_uses_language_factor_with_default_fallback(
        self,
    ) -> None:
        self.assertEqual(17, estimate_output_tokens(10, "ja", "en"))
        self.assertEqual(10, estimate_output_tokens(10, "xx", "yy"))

    def test_estimate_translation_costs_uses_model_pricing(self) -> None:
        input_cost, output_cost, total_cost = estimate_translation_costs(
            2000, 3000, "gpt-4o-mini"
        )

        self.assertAlmostEqual(0.0003, input_cost)
        self.assertAlmostEqual(0.0018, output_cost)
        self.assertAlmostEqual(0.0021, total_cost)


class TestAppHelperIntegrationTests(unittest.TestCase):
    def test_estimate_xlsx_file_reads_workbook_and_returns_cost_estimate(self) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet["A1"] = "Hello world"

        workbook_buffer = BytesIO()
        workbook.save(workbook_buffer)
        workbook_buffer.seek(0)

        input_tokens, output_tokens, total_cost = estimate_xlsx_file(
            workbook_buffer,
            "en",
            "fr",
            "gpt-4o-mini",
        )
        _, _, expected_total_cost = estimate_translation_costs(
            input_tokens,
            output_tokens,
            "gpt-4o-mini",
        )

        self.assertGreater(input_tokens, 0)
        self.assertEqual(
            estimate_output_tokens(input_tokens, "en", "fr"), output_tokens
        )
        self.assertAlmostEqual(expected_total_cost, total_cost)


class TestZipHelperUnitTests(unittest.TestCase):
    def test_is_zip_filename_accepts_zip_archives(self) -> None:
        self.assertTrue(is_zip_filename("archive.zip"))
        self.assertTrue(is_zip_filename("ARCHIVE.ZIP"))
        self.assertFalse(is_zip_filename("archive.xlsx"))

    def test_translated_zip_name_adds_target_language_suffix(self) -> None:
        self.assertEqual("source_en.zip", translated_zip_name("source.zip", "en"))
        self.assertEqual(
            "report.v1_fr.zip", translated_zip_name("/tmp/report.v1.zip", "fr")
        )

    def test_is_xlsx_filename_accepts_workbooks_and_filters_metadata(self) -> None:
        self.assertTrue(is_xlsx_filename("source.xlsx"))
        self.assertTrue(is_xlsx_filename("SOURCE.XLSX"))
        self.assertFalse(is_xlsx_filename("notes.txt"))
        self.assertFalse(is_xlsx_filename("~$source.xlsx"))
        self.assertFalse(is_xlsx_filename("._source.xlsx"))

    def test_safe_extract_zip_extracts_nested_archive_members(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = os.path.join(temp_dir, "source.zip")
            extract_dir = os.path.join(temp_dir, "extract")

            with zipfile.ZipFile(zip_path, "w") as archive:
                archive.writestr("nested/source.xlsx", b"workbook")
                archive.writestr("notes.txt", b"notes")

            safe_extract_zip(zip_path, extract_dir)

            self.assertEqual(
                b"workbook",
                Path(extract_dir, "nested", "source.xlsx").read_bytes(),
            )
            self.assertEqual(b"notes", Path(extract_dir, "notes.txt").read_bytes())

    def test_safe_extract_zip_rejects_unsafe_member_paths(self) -> None:
        unsafe_member_names = [
            "../evil.xlsx",
            "nested/../../evil.xlsx",
            "/tmp/evil.xlsx",
            "C:/tmp/evil.xlsx",
            "..\\evil.xlsx",
        ]

        for member_name in unsafe_member_names:
            with self.subTest(member_name=member_name):
                with tempfile.TemporaryDirectory() as temp_dir:
                    zip_path = os.path.join(temp_dir, "source.zip")
                    extract_dir = os.path.join(temp_dir, "extract")

                    with zipfile.ZipFile(zip_path, "w") as archive:
                        archive.writestr(member_name, b"evil")

                    with self.assertRaises(UnsafeZipError):
                        safe_extract_zip(zip_path, extract_dir)

                    self.assertFalse(Path(temp_dir, "evil.xlsx").exists())

    def test_find_xlsx_files_discovers_supported_workbooks_recursively(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            nested = root / "nested"
            nested.mkdir()
            keep_root = root / "root.xlsx"
            keep_nested = nested / "nested.XLSX"
            ignored_lock = root / "~$root.xlsx"
            ignored_resource = nested / "._nested.xlsx"
            ignored_text = root / "notes.txt"

            for path in [
                keep_root,
                keep_nested,
                ignored_lock,
                ignored_resource,
                ignored_text,
            ]:
                path.write_bytes(b"content")

            self.assertEqual(
                sorted([str(keep_root), str(keep_nested)]),
                find_xlsx_files(temp_dir),
            )


class TestDownloadArtifactServices(unittest.TestCase):
    def test_ensure_translation_work_dirs_creates_configured_dirs(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            work_dir = Path(temp_dir) / "work"
            tmp_dir = work_dir / "tmp"
            original_temp_dir = tempfile.tempdir
            tempfile.tempdir = "/tmp"

            try:
                with patch.dict(
                    os.environ,
                    {
                        TRANSLATION_WORK_DIR_ENV: str(work_dir),
                        TMPDIR_ENV: str(tmp_dir),
                    },
                ):
                    artifact_root = ensure_translation_work_dirs()

                self.assertEqual(work_dir, artifact_root)
                self.assertTrue(work_dir.is_dir())
                self.assertTrue(tmp_dir.is_dir())
                self.assertIsNone(tempfile.tempdir)
            finally:
                tempfile.tempdir = original_temp_dir

    def test_stage_download_artifact_writes_file_and_cleanup_removes_directory(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            session_state: dict[str, object] = {}
            artifact = stage_download_artifact(
                data=b"translated",
                filename="result.xlsx",
                session_state=session_state,
                work_dir=Path(temp_dir),
            )

            self.assertEqual(b"translated", artifact.path.read_bytes())
            self.assertTrue(artifact.directory.exists())
            self.assertEqual(
                str(artifact.directory),
                session_state[DOWNLOAD_ARTIFACT_DIR_KEY],
            )

            cleanup_download_artifact(session_state)

            self.assertFalse(artifact.directory.exists())
            self.assertNotIn(DOWNLOAD_ARTIFACT_DIR_KEY, session_state)

    def test_stage_download_artifact_removes_previous_session_artifact(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            session_state: dict[str, object] = {}
            previous_artifact = stage_download_artifact(
                data=b"first",
                filename="first.xlsx",
                session_state=session_state,
                work_dir=Path(temp_dir),
            )

            next_artifact = stage_download_artifact(
                data=b"second",
                filename="second.xlsx",
                session_state=session_state,
                work_dir=Path(temp_dir),
            )

            self.assertFalse(previous_artifact.directory.exists())
            self.assertTrue(next_artifact.directory.exists())


class TestActualUsageStateServices(unittest.TestCase):
    def test_source_change_clears_stored_actual_usage(self) -> None:
        session_state: dict[str, object] = {
            ACTUAL_USAGE_SOURCE_KEY: "source-a",
            ACTUAL_INPUT_TOKENS_KEY: 12,
            ACTUAL_OUTPUT_TOKENS_KEY: 18,
            ACTUAL_COST_KEY: 0.34,
            LEGACY_SHOW_ACTUAL_USAGE_KEY: True,
        }

        clear_actual_usage_on_source_change(session_state, "source-b")

        self.assertEqual("source-b", session_state[ACTUAL_USAGE_SOURCE_KEY])
        self.assertNotIn(ACTUAL_INPUT_TOKENS_KEY, session_state)
        self.assertNotIn(ACTUAL_OUTPUT_TOKENS_KEY, session_state)
        self.assertNotIn(ACTUAL_COST_KEY, session_state)
        self.assertNotIn(LEGACY_SHOW_ACTUAL_USAGE_KEY, session_state)

    def test_same_source_preserves_stored_actual_usage(self) -> None:
        session_state: dict[str, object] = {}
        store_actual_usage(
            session_state,
            "source-a",
            TranslationUsage(input_tokens=12, output_tokens=18, cost=0.34),
        )

        clear_actual_usage_on_source_change(session_state, "source-a")

        actual_usage = get_actual_usage_for_source(session_state, "source-a")
        self.assertIsNotNone(actual_usage)
        assert actual_usage is not None
        self.assertEqual(12, actual_usage.input_tokens)
        self.assertEqual(18, actual_usage.output_tokens)
        self.assertEqual(0.34, actual_usage.cost)

    def test_download_cleanup_does_not_clear_stored_actual_usage(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            session_state: dict[str, object] = {}
            stage_download_artifact(
                data=b"translated",
                filename="result.xlsx",
                session_state=session_state,
                work_dir=Path(temp_dir),
            )
            store_actual_usage(
                session_state,
                "source-a",
                TranslationUsage(input_tokens=12, output_tokens=18, cost=0.34),
            )

            cleanup_download_artifact(session_state)

            actual_usage = get_actual_usage_for_source(session_state, "source-a")
            self.assertIsNotNone(actual_usage)
            assert actual_usage is not None
            self.assertEqual(12, actual_usage.input_tokens)
            self.assertEqual(18, actual_usage.output_tokens)
            self.assertEqual(0.34, actual_usage.cost)
            self.assertNotIn(DOWNLOAD_ARTIFACT_DIR_KEY, session_state)


class FakeTranslator:
    """Fake translator that records constructor inputs and writes test output."""

    created: ClassVar[list["FakeTranslator"]] = []

    def __init__(
        self,
        input_path: str,
        target_language: str,
        model_name: str,
        max_workers: int,
        rpm_limit: int,
    ) -> None:
        """Create a fake translator instance and record constructor inputs."""
        self.input_path = input_path
        self.target_language = target_language
        self.model_name = model_name
        self.max_workers = max_workers
        self.rpm_limit = rpm_limit
        self.error: Optional[str] = None
        self.actual_input_tokens = 12
        self.actual_output_tokens = 18
        self.actual_cost = 0.34
        with open(self.input_path, "rb") as input_file:
            self.input_bytes = input_file.read()
        self.get_result_called = False
        FakeTranslator.created.append(self)

    def translate_with_progress(self) -> Iterable[float]:
        """Return deterministic progress values for orchestration tests."""
        return [0.25, 1.0]

    def get_result(self, output_path: str) -> Optional[str]:
        """Write a translated output file and return its path."""
        self.get_result_called = True
        with open(output_path, "wb") as output_file:
            output_file.write(b"translated workbook")
        return output_path


class ErrorTranslator(FakeTranslator):
    """Fake translator that reports a translation error."""

    def __init__(
        self,
        input_path: str,
        target_language: str,
        model_name: str,
        max_workers: int,
        rpm_limit: int,
    ) -> None:
        """Create a fake translator that reports an error."""
        super().__init__(
            input_path=input_path,
            target_language=target_language,
            model_name=model_name,
            max_workers=max_workers,
            rpm_limit=rpm_limit,
        )
        self.error: Optional[str] = "boom"


class MissingResultTranslator(FakeTranslator):
    """Fake translator that cannot produce an output file."""

    def get_result(self, output_path: str) -> Optional[str]:
        """Return no output path to simulate a failed result write."""
        self.get_result_called = True
        return None


class TestSingleXLSXTranslationOrchestration(unittest.TestCase):
    def setUp(self) -> None:
        FakeTranslator.created = []

    def test_translated_xlsx_name_adds_target_language_suffix(self) -> None:
        self.assertEqual("source_en.xlsx", translated_xlsx_name("source.xlsx", "en"))
        self.assertEqual(
            "report.v1_fr.xlsx", translated_xlsx_name("/tmp/report.v1.xlsx", "fr")
        )

    def test_translate_single_xlsx_returns_download_payload_and_usage(self) -> None:
        progress_values: list[float] = []

        result = translate_single_xlsx(
            uploaded_bytes=b"source workbook",
            original_filename="source.xlsx",
            target_language="en",
            model_name="gpt-4o-mini",
            max_workers=3,
            rpm_limit=99,
            progress_callback=progress_values.append,
            translator_factory=FakeTranslator,
        )

        self.assertEqual(b"translated workbook", result.data)
        self.assertEqual("source_en.xlsx", result.filename)
        self.assertEqual(XLSX_MIME_TYPE, result.mime_type)
        self.assertEqual(12, result.usage.input_tokens)
        self.assertEqual(18, result.usage.output_tokens)
        self.assertEqual(0.34, result.usage.cost)
        self.assertEqual([0.25, 1.0], progress_values)

        fake = FakeTranslator.created[0]
        self.assertEqual(b"source workbook", fake.input_bytes)
        self.assertEqual("en", fake.target_language)
        self.assertEqual("gpt-4o-mini", fake.model_name)
        self.assertEqual(3, fake.max_workers)
        self.assertEqual(99, fake.rpm_limit)
        self.assertTrue(fake.get_result_called)

    def test_translate_single_xlsx_raises_when_translation_reports_error(self) -> None:
        with self.assertRaisesRegex(TranslationError, "boom"):
            translate_single_xlsx(
                uploaded_bytes=b"source workbook",
                original_filename="source.xlsx",
                target_language="en",
                model_name="gpt-4o-mini",
                max_workers=3,
                translator_factory=ErrorTranslator,
            )

    def test_translate_single_xlsx_raises_when_result_is_missing(self) -> None:
        with self.assertRaisesRegex(TranslationError, "Translation failed"):
            translate_single_xlsx(
                uploaded_bytes=b"source workbook",
                original_filename="source.xlsx",
                target_language="en",
                model_name="gpt-4o-mini",
                max_workers=3,
                translator_factory=MissingResultTranslator,
            )


class TestZipXLSXTranslationOrchestration(unittest.TestCase):
    def setUp(self) -> None:
        FakeTranslator.created = []

    def create_zip_bytes(self, members: dict[str, bytes]) -> bytes:
        """Create an in-memory zip archive from member bytes."""
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            for member_name, member_bytes in members.items():
                archive.writestr(member_name, member_bytes)
        return buffer.getvalue()

    def test_translate_xlsx_zip_returns_translated_archive_and_usage(self) -> None:
        progress_values: list[ZipTranslationProgress] = []
        archive_bytes = self.create_zip_bytes(
            {
                "root.xlsx": b"root workbook",
                "nested/team.xlsx": b"nested workbook",
                "notes.txt": b"ignored",
                "~$lock.xlsx": b"ignored lock",
            }
        )

        result = translate_xlsx_zip(
            uploaded_bytes=archive_bytes,
            original_filename="archive.zip",
            target_language="en",
            model_name="gpt-4o-mini",
            max_workers=3,
            rpm_limit=99,
            progress_callback=progress_values.append,
            translator_factory=FakeTranslator,
        )

        self.assertEqual("archive_en.zip", result.filename)
        self.assertEqual(ZIP_MIME_TYPE, result.mime_type)
        self.assertEqual(2, result.workbook_count)
        self.assertEqual(24, result.usage.input_tokens)
        self.assertEqual(36, result.usage.output_tokens)
        self.assertAlmostEqual(0.68, result.usage.cost)

        with zipfile.ZipFile(BytesIO(result.data)) as archive:
            self.assertEqual(
                ["nested/team_en.xlsx", "root_en.xlsx"],
                sorted(archive.namelist()),
            )
            self.assertEqual(
                b"translated workbook",
                archive.read("nested/team_en.xlsx"),
            )
            self.assertEqual(b"translated workbook", archive.read("root_en.xlsx"))

        self.assertEqual(2, len(FakeTranslator.created))
        self.assertEqual("en", FakeTranslator.created[0].target_language)
        self.assertEqual("gpt-4o-mini", FakeTranslator.created[0].model_name)
        self.assertEqual(3, FakeTranslator.created[0].max_workers)
        self.assertEqual(99, FakeTranslator.created[0].rpm_limit)
        self.assertEqual(
            ["nested/team.xlsx", "root.xlsx"],
            [
                progress.relative_path
                for progress in progress_values
                if progress.workbook_progress == 1.0
            ],
        )
        self.assertEqual(
            [0.5, 1.0],
            [
                progress.overall_progress
                for progress in progress_values
                if progress.workbook_progress == 1.0
            ],
        )

    def test_translate_xlsx_zip_raises_when_archive_has_no_workbooks(self) -> None:
        archive_bytes = self.create_zip_bytes({"notes.txt": b"notes"})

        with self.assertRaisesRegex(TranslationError, "No .xlsx workbooks"):
            translate_xlsx_zip(
                uploaded_bytes=archive_bytes,
                original_filename="archive.zip",
                target_language="en",
                model_name="gpt-4o-mini",
                max_workers=3,
                translator_factory=FakeTranslator,
            )

    def test_translate_xlsx_zip_includes_relative_path_in_translation_errors(
        self,
    ) -> None:
        archive_bytes = self.create_zip_bytes(
            {"nested/source.xlsx": b"source workbook"}
        )

        with self.assertRaisesRegex(
            TranslationError,
            "nested/source.xlsx: boom",
        ):
            translate_xlsx_zip(
                uploaded_bytes=archive_bytes,
                original_filename="archive.zip",
                target_language="en",
                model_name="gpt-4o-mini",
                max_workers=3,
                translator_factory=ErrorTranslator,
            )

    def test_translate_xlsx_zip_includes_relative_path_when_result_is_missing(
        self,
    ) -> None:
        archive_bytes = self.create_zip_bytes(
            {"nested/source.xlsx": b"source workbook"}
        )

        with self.assertRaisesRegex(
            TranslationError,
            "Translation failed for nested/source.xlsx",
        ):
            translate_xlsx_zip(
                uploaded_bytes=archive_bytes,
                original_filename="archive.zip",
                target_language="en",
                model_name="gpt-4o-mini",
                max_workers=3,
                translator_factory=MissingResultTranslator,
            )


if __name__ == "__main__":
    unittest.main()
